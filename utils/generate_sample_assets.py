"""
Script to generate sample PDF, DOCX, and PPTX test files in data/sample_materials/
"""

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config


def generate_sample_files():
    notes_txt_path = config.SAMPLE_DIR / "Machine_Learning_Viva_Notes.txt"
    if not notes_txt_path.exists():
        print("Sample txt file not found.")
        return

    with open(notes_txt_path, "r", encoding="utf-8") as f:
        content = f.read()

    # 1. Generate DOCX
    import docx

    doc = docx.Document()
    doc.add_heading("Machine Learning & Algorithms Viva Notes", level=0)

    for line in content.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line.replace("# ", ""), level=1)
        elif line.startswith("## "):
            doc.add_heading(line.replace("## ", ""), level=2)
        elif line.strip():
            doc.add_paragraph(line.strip())

    docx_path = config.SAMPLE_DIR / "Machine_Learning_Viva_Notes.docx"
    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path}")

    # 2. Generate PPTX
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    # Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Machine Learning Viva Preparation"
    subtitle.text = "Key Concepts, Random Forest, CNN & Graph Traversals"

    # Content slides
    sections = content.split("\n# ")
    for sec in sections:
        if not sec.strip():
            continue
        lines = sec.strip().split("\n")
        sec_title = lines[0].replace("#", "").strip()

        content_layout = prs.slide_layouts[1]
        c_slide = prs.slides.add_slide(content_layout)
        c_slide.shapes.title.text = sec_title[:45]
        tf = c_slide.placeholders[1].text_frame

        body_lines = [l for l in lines[1:8] if l.strip() and not l.startswith("##")]
        for idx, bline in enumerate(body_lines[:5]):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = bline.replace("-", "").strip()[:100]
            p.font.size = Pt(16)

    pptx_path = config.SAMPLE_DIR / "Machine_Learning_Viva_Slides.pptx"
    prs.save(str(pptx_path))
    print(f"Generated PPTX: {pptx_path}")

    # 3. Generate PDF using PyMuPDF (fitz)
    import fitz

    pdf_doc = fitz.open()
    # Create pages from text
    lines = content.split("\n")
    page_lines = []
    page_num = 1

    for line in lines:
        page_lines.append(line)
        if len(page_lines) >= 35:
            page = pdf_doc.new_page()
            text_str = "\n".join(page_lines)
            page.insert_text(fitz.Point(40, 50), text_str, fontsize=10)
            page_lines = []
            page_num += 1

    if page_lines:
        page = pdf_doc.new_page()
        text_str = "\n".join(page_lines)
        page.insert_text(fitz.Point(40, 50), text_str, fontsize=10)

    pdf_path = config.SAMPLE_DIR / "Machine_Learning_Viva_Notes.pdf"
    pdf_doc.save(str(pdf_path))
    pdf_doc.close()
    print(f"Generated PDF: {pdf_path}")


if __name__ == "__main__":
    generate_sample_files()
